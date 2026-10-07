import pandas as pd
import numpy as np
import sys
from pathlib import Path
sys.path.append('../../src')
from data_imports import *

import datetime
import numbers
import shutil
from openpyxl import load_workbook

def archive_workbook(path, archive_dir=None):
    '''
    Copy the workbook to <archive_dir>/<name>_<YYYYMMDD-HHMMSS>.xlsx (default archive_dir: an "archive" folder next to it).
    Returns the path of the archived copy.
    '''
    path = pathlib.Path(path)
    archive_dir = pathlib.Path(archive_dir) if archive_dir else path.parent / "archive"
    archive_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = archive_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, dest)
    return dest

def _cell_value(v):
    '''
    Convert a DataFrame value to an openpyxl cell value. Returns (value, is_text).
    Missing values become empty cells; booleans and numbers keep their type; everything else is written as text.
    '''
    if v is None or v is pd.NA or v is pd.NaT or (isinstance(v, float) and np.isnan(v)):
        return None, False
    if isinstance(v, (bool, np.bool_)):
        return bool(v), False
    if isinstance(v, numbers.Integral):
        return int(v), False
    if isinstance(v, numbers.Real):
        return float(v), False
    return str(v), True

def write_supplementary_tables(tables, path=SUPPLEMENTARY_TABLES_PATH, archive=True, archive_dir=None):
    '''
    Replace the contents of existing sheets in the Supplementary Tables workbook; all other sheets are left untouched.
    tables: {sheet name: DataFrame}. A named index (eg. biosample_id) is written as the first column.
    Strings are written as text cells ('@' format), so Excel never reinterprets them (eg. '7316-3' as a date).
    archive: copy the existing workbook to a timestamped file first.
    '''
    path = pathlib.Path(path)
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        raise RuntimeError(f"{path.name} appears to be open in Excel ({lock.name} exists); close it first.")
    wb = load_workbook(path)
    missing = [sheet for sheet in tables if sheet not in wb.sheetnames]
    if missing:
        raise KeyError(f"Sheets not found in {path.name}: {missing}")
    if archive:
        print(f"Archived previous version to {archive_workbook(path, archive_dir)}")

    for sheet, df in tables.items():
        df = df.reset_index() if df.index.name else df
        ws = wb[sheet]
        ws.delete_rows(1, ws.max_row)
        for c, name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=c, value=str(name))
            cell.number_format = "@"
        for r, row in enumerate(df.itertuples(index=False), start=2):
            for c, v in enumerate(row, start=1):
                value, is_text = _cell_value(v)
                cell = ws.cell(row=r, column=c, value=value)
                if is_text:
                    cell.number_format = "@"
        print(f"Wrote {len(df)} rows x {len(df.columns)} columns to '{sheet}'")
    wb.save(path)